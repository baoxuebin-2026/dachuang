from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "大创项目纯气体方向会议汇报稿_2026-10-24.docx"


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=11, bold=None):
    run.font.name = latin
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), latin)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_body(doc, text, first_line=True, bold_lead=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.45
    pf.space_after = Pt(6)
    if first_line:
        pf.first_line_indent = Pt(22)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=11, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=11)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, east_asia="黑体", latin="Arial", size=14 if level == 1 else 12, bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.font.size = Pt(11)

title_style = styles["Title"]
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_ppr = title_style.element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for name in ("Heading 1", "Heading 2"):
    style = styles[name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    style.font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
r = title.add_run("大创项目纯气体方向会议汇报稿")
set_run_font(r, east_asia="黑体", latin="Arial", size=20, bold=True)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(4)
r = subtitle.add_run("面向山西焦化园区的跨设备低浓度气体识别")
set_run_font(r, east_asia="黑体", latin="Arial", size=12, bold=True)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_after = Pt(14)
r = meta.add_run("汇报日期：2026年10月24日")
set_run_font(r, size=10.5)

add_heading(doc, "一 先说结论", 1)
add_body(
    doc,
    "目前纯气体方向已经形成了一条可以继续写论文的主线：模型只使用其他设备采集的较高浓度数据训练，在完全留出一台新设备的情况下，识别这台设备上的较低浓度气体。我们没有发明一种全新的分类器，而是把分级幅值缩放这种简单的数据增强方法，用到了“换设备加浓度下降”同时发生的问题上。主数据集结果比较稳定，两套补充数据则说明了方法在长期漂移、湿度和加热变化下的适用边界。",
)
add_body(
    doc,
    "一句话概括当前工作：先把高浓度训练信号按0.50、0.75和1.00倍缩放，让模型提前见到较弱的响应，再拿完全没有参与训练的新设备低浓度数据进行测试。",
    bold_lead="一句话概括当前工作：",
)

add_heading(doc, "二 这项研究到底在解决什么", 1)
add_body(
    doc,
    "实际使用气体监测设备时，同型号传感器也会有个体差异。一台设备上训练好的模型，换到另一台设备后不一定还能用。如果新设备检测到的气体浓度又更低，传感器响应会进一步变弱，模型就更容易认错。我们的研究问题可以直接表述为：人工智能只看过其他设备的高浓度气体数据，换一台新设备以后，它还能不能认出低浓度气体？",
)
add_body(
    doc,
    "现阶段研究对象是CO、甲烷、乙醇和乙烯四类气体。焦化园区是应用动机，因为CO和可燃、有毒气体监测确实存在设备更换、标定样本不足和低浓度识别问题；实验本身使用公开受控数据，重点验证算法环节。",
)

add_heading(doc, "三 主实验已经完成到什么程度", 1)
add_body(
    doc,
    "主实验采用UCI 361 Twin gas sensor arrays数据。数据由5台结构相同的电子鼻设备采集，每台设备有8个气敏通道，包含4类气体、10个浓度档和640次有效供气记录。每轮完整留出1台设备作为测试设备，其余4台作为训练设备；训练只使用第4至第10浓度档，测试只使用留出设备的第1至第3浓度档。这样训练设备和测试设备不同，高低浓度范围也不重叠。5轮合计得到192次低浓度测试。",
)
add_body(
    doc,
    "处理时先从供气前后的八通道对数电阻中提取20乘8的动态响应矩阵，再只对训练数据做0.50、0.75和1.00倍缩放。为了判断提升是不是来自分类器本身，实验同时使用SVM和一维CNN，并与原始响应、逐试次RMS归一化、平滑噪声和时间平移进行比较。",
)
add_body(
    doc,
    "结果方面，SVM的宏平均F1由0.736提高到0.963；一维CNN由0.691提高到0.974。增强CNN在192次测试中正确识别187次，48次CO全部识别正确；剩余错误主要是5次非CO被误报为CO。因此，当前方法解决的是整体四气体识别，同时保住了CO召回率，但误报仍然需要在后续应用中考虑。",
)
add_body(
    doc,
    "我们还测试了少量目标设备标定。单独加入20次目标设备高浓度样本并没有明显改善；把这20次样本与幅值增强结合后，SVM在不同抽样下达到0.968至0.990的宏平均F1，正确识别186至190次。这说明少量标定主要补设备差异，幅值增强主要补浓度范围，两者作用不同，可以组合。",
)

add_heading(doc, "四 另外两套数据为什么效果不高", 1)
add_body(
    doc,
    "另外两套数据不是简单的“格式不一样所以跑不好”。更关键的是，它们里面发生的变化已经不只是浓度下降。我们的幅值增强相当于把信号整体调小，适合处理“响应形状大致不变、主要是幅度减弱”的情况；如果传感器的不同通道、湿度背景或加热周期一起变化，单纯缩放幅值就不够了。",
)

table = doc.add_table(rows=1, cols=4)
table.autofit = False
widths = [Cm(2.1), Cm(4.1), Cm(3.4), Cm(5.1)]
headers = ["数据集", "实际检验的问题", "实验结果", "为什么不能照搬主实验结论"]
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.width = widths[i]
    set_cell_shading(cell, "D9E6F2")
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, east_asia="黑体", latin="Arial", size=9.5, bold=True)

rows = [
    (
        "UCI 361 主数据",
        "5台设备之间迁移，同时从高浓度训练到低浓度测试，识别4类气体",
        "SVM F1：0.736→0.963；CNN F1：0.691→0.974",
        "与方法假设最匹配，是论文的主要证据",
    ),
    (
        "UCI 270",
        "同一套阵列跨约20个月以上老化过程，从早期高浓度识别后期低浓度6类气体",
        "SVM F1：0.225→0.452，5个测试批次均提高",
        "无CO、无多台完整设备；只有128维人工特征；长期老化使各通道发生非一致漂移，不是整体缩放能解决的",
    ),
    (
        "UCI 487",
        "7只同型号传感器跨日期识别低浓度CO和空气",
        "SVM F1：0.529→0.593；15/21个评价块提高；CO召回率0.512",
        "低浓度CO与湿度、日期漂移和周期加热同时变化；只有CO/空气二分类，无法检查其他气体误报为CO",
    ),
]
for row_idx, values in enumerate(rows, start=1):
    cells = table.add_row().cells
    for i, text in enumerate(values):
        cells[i].width = widths[i]
        set_cell_margins(cells[i])
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if row_idx % 2 == 0:
            set_cell_shading(cells[i], "F5F7FA")
        p = cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        set_run_font(r, size=9.2, bold=(i == 0))

set_table_borders(table)
table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

add_body(
    doc,
    "所以，这两套补充数据不能说成完全失败。两套数据上幅值增强都有提升，说明“补充弱响应”这个方向不是只对主数据有效；但绝对分数仍然不高，说明它不能同时解决长期老化、湿度变化和加热相位变化。它们更适合放在论文里说明方法边界，而不是包装成第二次、第三次成功复现。",
)

add_heading(doc, "五 当前可以怎样概括创新点", 1)
add_body(
    doc,
    "当前创新点属于应用方法创新，不是宣称国内外从未出现过的全新算法。具体来说，是把分级幅值缩放用于跨设备、跨浓度的多气体识别，并采用整台设备留出、高低浓度不重叠的实验方式进行验证。与只随机划分样本相比，这种设置更接近设备更换后的真实困难；同时用SVM和CNN得到一致趋势，说明提升主要来自训练数据覆盖了弱响应，而不是某一种网络结构偶然有效。",
)
add_body(
    doc,
    "从本科生论文和《山西电子技术》这类应用型期刊的尺度看，这条创新主线已经具备写作基础。论文最有价值的不是网络有多复杂，而是问题设置明确、实验划分严格、主结果提升明显，并且把补充数据上的负面结果如实解释清楚。",
)

add_heading(doc, "六 当前论文和材料进度", 1)
add_bullet(doc, "中文投稿初稿已经完成，标题暂定为《面向山西焦化园区的跨设备低浓度气体识别方法》。")
add_bullet(doc, "正文已经整理为8页Word稿，包含中文和英文摘要、4张结果图、4个表格和15条参考文献。")
add_bullet(doc, "四张结果图已经改为常见的柱状图和混淆矩阵，并同时保留PNG、SVG和PDF格式及Python绘图代码。")
add_bullet(doc, "主数据、补充压力测试、少量目标设备标定和失败路线都已经保留，没有把负面结果删掉或改写成成功。")
add_bullet(doc, "论文相关Word、图表和脚本已经提交到项目仓库的研究分支。")

add_heading(doc, "七 下一步安排", 1)
add_body(
    doc,
    "下一步不再盲目增加数据集，而是先把当前稿件做成可以投稿的版本。具体需要核对作者和基金信息，统一图表、术语和参考文献格式，检查所有数字是否能够从实验结果文件中复现，再按目标期刊模板压缩篇幅。实验上只补真正影响结论的内容，例如按设备统计波动、给出必要的显著性或稳定性说明，以及进一步解释CO误报。",
)
add_body(
    doc,
    "如果以后能够获得带湿度、混合气干扰和多台完整设备的新数据，再考虑加入湿度补偿或通道级校正；目前不应该为了显得工作量大，把任务完全不同的数据强行合并训练。",
)

add_heading(doc, "八 汇报结束语", 1)
add_body(
    doc,
    "以上是目前纯气体方向的进展。简单总结就是：主数据已经证明，分级幅值增强能够明显改善换设备后的低浓度四气体识别；两套补充数据说明，遇到长期漂移、湿度和加热变化时，单纯缩放幅值只能部分起作用。当前最合理的工作重点是把主线写清楚、把适用边界说明白，并完成投稿稿件，而不是继续无目的地堆数据集。",
)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("大创项目纯气体方向会议汇报稿")
set_run_font(fr, size=9)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
