"""Apply journal-draft typography and layout to the current manuscript DOCX."""

from __future__ import annotations

from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_run_font(run, *, east_asia="宋体", latin="Times New Roman", size=10.5,
                 bold=None):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def format_paragraph(paragraph):
    text = paragraph.text.strip()
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(4)
    fmt.line_spacing = 1.25

    if paragraph.style.name in {"Heading 2", "Heading 3"}:
        level = paragraph.style.name
        fmt.space_before = Pt(8 if level == "Heading 2" else 6)
        fmt.space_after = Pt(4)
        fmt.keep_with_next = True
        fmt.first_line_indent = None
        for run in paragraph.runs:
            set_run_font(run, east_asia="黑体", size=12 if level == "Heading 2" else 11,
                         bold=True)
        return

    if text.startswith("图") and len(text) < 80:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.keep_with_next = False
        fmt.space_before = Pt(2)
        fmt.space_after = Pt(6)
        fmt.first_line_indent = None
        for run in paragraph.runs:
            set_run_font(run, size=9)
        return

    if text.startswith("表") and len(text) < 80:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.keep_with_next = True
        fmt.space_before = Pt(5)
        fmt.space_after = Pt(3)
        fmt.first_line_indent = None
        for run in paragraph.runs:
            set_run_font(run, size=9, bold=True)
        return

    if paragraph._p.xpath(".//w:drawing"):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.keep_with_next = True
        fmt.space_before = Pt(3)
        fmt.space_after = Pt(1)
        fmt.first_line_indent = None
        return

    if text:
        fmt.first_line_indent = Cm(0.74)
    for run in paragraph.runs:
        set_run_font(run, size=10.5)


def main(src: Path, dst: Path):
    doc = Document(src)
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    # Main title.
    title = doc.paragraphs[0]
    title.style = doc.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    for run in title.runs:
        set_run_font(run, east_asia="黑体", size=18, bold=True)

    # Chinese author and affiliation.
    for idx in (1, 2):
        p = doc.paragraphs[idx]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.space_after = Pt(5 if idx == 1 else 10)
        p.paragraph_format.keep_with_next = True
        for run in p.runs:
            set_run_font(run, size=10.5)

    for p in doc.paragraphs[3:]:
        format_paragraph(p)

    # Compact GB/T-style references with a hanging indent so the last page is
    # not left mostly empty.
    in_references = False
    for p in doc.paragraphs:
        if p.text.strip() == "参考文献":
            in_references = True
            continue
        if in_references and p.text.strip():
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.first_line_indent = Cm(-0.74)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                set_run_font(run, size=9)

    # Front matter uses no first-line indentation.
    for idx in range(3, 11):
        p = doc.paragraphs[idx]
        p.paragraph_format.first_line_indent = None
    for idx in (6, 7, 8):
        doc.paragraphs[idx].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[6].paragraph_format.space_before = Pt(8)
    doc.paragraphs[6].paragraph_format.keep_with_next = True
    for run in doc.paragraphs[6].runs:
        set_run_font(run, size=12, bold=True)

    # Tables: compact but readable, with repeating header and visible borders.
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_table_borders(table)
        header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header_tr_pr.append(repeat)
        for row_idx, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell)
                if row_idx == 0:
                    set_cell_shading(cell, "DCE6F1")
                elif row_idx % 2 == 0:
                    set_cell_shading(cell, "F7F9FB")
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = None
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.05
                    if row_idx == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        set_run_font(run, size=8.5, bold=True if row_idx == 0 else None)

    # Enlarge charts to a comfortable journal proof width.
    widths_cm = [14.2, 14.2, 10.7, 14.2]
    for shape, width_cm in zip(doc.inline_shapes, widths_cm):
        ratio = shape.height / shape.width
        shape.width = Cm(width_cm)
        shape.height = int(shape.width * ratio)

    # Do not split logical figure/table units across pages when avoidable.
    for p in doc.paragraphs:
        if p._p.xpath(".//w:drawing"):
            p.paragraph_format.keep_with_next = True

    doc.core_properties.title = "面向山西焦化园区的跨设备低浓度气体识别方法"
    doc.core_properties.subject = "投稿初稿"
    doc.save(dst)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: format_current_manuscript_docx.py INPUT.docx OUTPUT.docx")
    main(Path(sys.argv[1]), Path(sys.argv[2]))
