#!/usr/bin/env python3
"""Audit the published WATCH4ERS workbook without estimating alarm accuracy."""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


EXPECTED_SHA256 = "331677b23555e58dc9e081e7f13fe0b18a7915b7194a95cc53db37e8fe872730"


def run(source: Path) -> dict:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if digest != EXPECTED_SHA256:
        raise ValueError(f"WATCH4ERS original workbook SHA-256 mismatch: {digest}")

    book = load_workbook(source, read_only=True, data_only=True)
    if len(book.sheetnames) != 1:
        raise ValueError(f"Expected exactly one worksheet: {book.sheetnames}")
    sheet = book.active
    data = list(sheet.values)
    expected_header = (
        " Time", "TDLAS CH4 concentration (ppm)",
        "NDIR CH4 concentration (ppm)", "MOx CH4 concentration (ppm)", "Gas on"
    )
    if tuple(data[0]) != expected_header:
        raise ValueError(f"Unexpected variable names: {data[0]}")
    if any(len(row) != 5 or any(v is None for v in row) for row in data[1:]):
        raise ValueError("Workbook has missing fields")

    x = np.asarray(data[1:], dtype=float)
    if not np.isfinite(x).all() or (np.diff(x[:, 0]) <= 0).any():
        raise ValueError("Missing or nonmonotone timestamp or sensor value")
    gas = x[:, 4]
    if not set(np.unique(gas)) <= {0, 10}:
        raise ValueError("Undocumented gas switch values")
    starts = np.where((gas == 10) & np.r_[True, gas[:-1] != 10])[0]
    ends = np.where((gas == 0) & np.r_[False, gas[:-1] == 10])[0]
    seconds = (x[:, 0] - x[0, 0]) * 86400
    if len(starts) != 15 or len(ends) != 15:
        raise ValueError(f"Expected 15 valve pulses, found {len(starts)} starts / {len(ends)} ends")

    # The workbook has no experiment-ID column; published Table 1 establishes
    # five experiments of three pulses. Do not infer a reset from clock data.
    by_triple = [starts[i:i + 3].tolist() for i in range(0, 15, 3)]
    intervals = np.diff(seconds)
    return {
        "source": str(source),
        "source_sha256": digest,
        "workbook_sheets": book.sheetnames,
        "data_rows": int(len(x)),
        "elapsed_s": round(float(seconds[-1]), 3),
        "sample_interval_s": {
            "min": round(float(intervals.min()), 3),
            "median": round(float(np.median(intervals)), 3),
            "max": round(float(intervals.max()), 3),
            "rounded_counts": dict(sorted(Counter(np.rint(intervals).astype(int).tolist()).items())),
        },
        "gas_switch_values": dict(sorted(Counter(gas.astype(int).tolist()).items())),
        "pulse_count": len(starts),
        "pulse_onsets_elapsed_s": [round(float(seconds[i]), 3) for i in starts],
        "pulse_offsets_elapsed_s": [round(float(seconds[i]), 3) for i in ends],
        "onsets_by_published_triple": [
            [round(float(seconds[i]), 3) for i in group] for group in by_triple
        ],
        "initial_clean_air_observations_before_first_pulse": int(starts[0]),
        "observed_experiment_id_column": False,
        "protocol_watch_alarm_evaluation": "stop_insufficient_first_group_baseline",
        "reason": (
            "First pulse follows only one recorded clean-air observation, whereas the "
            "frozen protocol requires at least three for a within-group threshold. "
            "Other groups are concatenated without an experiment-ID or reset marker; "
            "preceding gas-off rows may contain residual methane."
        ),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    report = run(args.source)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({k: report[k] for k in (
        "source_sha256", "data_rows", "pulse_count", "initial_clean_air_observations_before_first_pulse",
        "protocol_watch_alarm_evaluation")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
